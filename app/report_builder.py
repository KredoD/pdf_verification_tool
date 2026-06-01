from __future__ import annotations

import html
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.models import CheckReport, EvidenceItem
from app.verification import FACT_LABELS
from app.utils import one_line

STATUS_LABELS = {
    "correct": "Корректно",
    "possible_match": "Похоже, нужна проверка",
    "mismatch": "Расхождение",
    "missing_in_reference": "Нет в эталоне",
    "missing_in_document": "Нет в документе",
    "not_checked": "Не проверено",
}

STATUS_COLORS = {
    "correct": "D9EAD3",
    "possible_match": "FFF2CC",
    "mismatch": "F4CCCC",
    "missing_in_reference": "FCE5CD",
    "missing_in_document": "FCE5CD",
    "not_checked": "E7E6E6",
}


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _cell(cell, text: str, *, bold: bool = False, fill: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text or ""))
    run.bold = bold
    run.font.size = Pt(8)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    if fill:
        _shade(cell, fill)


def _evidence_text(e: EvidenceItem | None) -> str:
    if not e:
        return "—"
    return f"{e.source_file}; {e.location}; фрагмент: {one_line(e.quote, 260)}"


def write_json_report(report: CheckReport, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")


def write_docx_report(report: CheckReport, path: Path) -> None:
    """Human-readable DOCX for employee/client archive."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)

    title = doc.add_heading("Отчет первичной проверки документации для ВНИИИМТ", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Проект: {report.project_name}")
    doc.add_paragraph(f"Эталонный документ: {report.reference_file}")
    doc.add_paragraph(f"Дата формирования: {report.generated_at}")
    doc.add_paragraph(f"Итог: {report.summary.get('human_status', '')}")

    doc.add_heading("1. Что извлек скрипт из документов", level=1)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Документ", "Роль", "Символов текста", "OCR/картинки", "Извлеченные типы данных"]
    for i, h in enumerate(headers):
        _cell(t.rows[0].cells[i], h, bold=True, fill="E9EEF3")
    for d in report.documents:
        c = t.add_row().cells
        extracted = "; ".join(f"{FACT_LABELS.get(k, k)}: {len(v)}" for k, v in d.facts.items() if v) or "не найдено"
        ocr = []
        if d.stats.get("pdf_pages_ocr_used"):
            ocr.append(f"PDF OCR страниц: {d.stats.get('pdf_pages_ocr_used')}")
        if d.stats.get("docx_images_ocr_used"):
            ocr.append(f"DOCX OCR картинок: {d.stats.get('docx_images_ocr_used')}/{d.stats.get('docx_images_total')}")
        _cell(c[0], d.filename)
        _cell(c[1], d.role)
        _cell(c[2], str(d.stats.get("total_text_chars", 0)))
        _cell(c[3], "; ".join(ocr) or "не использовался")
        _cell(c[4], extracted)


    doc.add_heading("2. Проверочные артефакты по типу пакета", level=1)
    art_table = doc.add_table(rows=1, cols=5)
    art_table.style = "Table Grid"
    for i, h in enumerate(["Статус", "Правило", "Найденные файлы", "Что требуется", "Комментарий"]):
        _cell(art_table.rows[0].cells[i], h, bold=True, fill="E9EEF3")
    art_colors = {"found": "D9EAD3", "missing": "F4CCCC", "needs_review": "FFF2CC"}
    art_labels = {"found": "Найдено", "missing": "Не найдено", "needs_review": "Нужна проверка"}
    for req in report.artifact_requirements:
        c = art_table.add_row().cells
        _cell(c[0], art_labels.get(req.status, req.status), fill=art_colors.get(req.status, "E7E6E6"))
        _cell(c[1], req.title)
        _cell(c[2], "; ".join(req.found_files) or "—")
        _cell(c[3], req.expected)
        _cell(c[4], req.comment)

    if report.address_consistency:
        doc.add_heading("3. Отдельная сверка адресов", level=1)
        addr_table = doc.add_table(rows=1, cols=6)
        addr_table.style = "Table Grid"
        for i, h in enumerate(["Статус", "Документ", "Адрес в документе", "Эталон", "Совпадение", "Комментарий"]):
            _cell(addr_table.rows[0].cells[i], h, bold=True, fill="E9EEF3")
        for row in report.address_consistency:
            c = addr_table.add_row().cells
            _cell(c[0], STATUS_LABELS.get(row.status, row.status), fill=STATUS_COLORS.get(row.status, "FFFFFF"))
            _cell(c[1], row.left_file)
            _cell(c[2], one_line(row.left_address, 320))
            _cell(c[3], one_line(row.right_address, 320))
            _cell(c[4], f"{row.score}%")
            _cell(c[5], row.comment)

    doc.add_heading("4. Проверка соответствия эталону", level=1)
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Статус", "Поле", "Проверяемый документ", "Значение в документе", "Значение в эталоне", "Совпадение", "Откуда взято в документе", "Комментарий"]
    for i, h in enumerate(headers):
        _cell(table.rows[0].cells[i], h, bold=True, fill="E9EEF3")

    # Put problematic rows first.
    order = {"mismatch": 0, "missing_in_document": 1, "missing_in_reference": 2, "possible_match": 3, "correct": 4}
    for row in sorted(report.rows, key=lambda r: (order.get(r.status, 9), r.checked_file, r.fact_type)):
        c = table.add_row().cells
        fill = STATUS_COLORS.get(row.status, "FFFFFF")
        _cell(c[0], STATUS_LABELS.get(row.status, row.status), fill=fill)
        _cell(c[1], FACT_LABELS.get(row.fact_type, row.fact_type))
        _cell(c[2], row.checked_file)
        _cell(c[3], one_line(row.checked_value, 260))
        _cell(c[4], one_line(row.reference_value, 260))
        _cell(c[5], f"{row.score}%")
        _cell(c[6], _evidence_text(row.checked_evidence))
        _cell(c[7], row.comment)

    doc.add_heading("5. Доказательства по эталонному документу", level=1)
    ref_doc = next((d for d in report.documents if d.filename == report.reference_file), None)
    if ref_doc:
        ev_table = doc.add_table(rows=1, cols=5)
        ev_table.style = "Table Grid"
        for i, h in enumerate(["Поле", "Значение", "Файл", "Место", "Фрагмент"]):
            _cell(ev_table.rows[0].cells[i], h, bold=True, fill="E9EEF3")
        for fact_type, items in ref_doc.facts.items():
            for item in items[:80]:
                c = ev_table.add_row().cells
                _cell(c[0], FACT_LABELS.get(fact_type, fact_type))
                _cell(c[1], item.value)
                _cell(c[2], item.source_file)
                _cell(c[3], item.location)
                _cell(c[4], one_line(item.quote, 350))

    doc.save(path)


def html_report(report: CheckReport) -> str:
    """Browser report shown automatically after processing."""
    status_counts = report.summary.get("status_counts", {})
    rows_html = []
    order = {"mismatch": 0, "missing_in_document": 1, "missing_in_reference": 2, "possible_match": 3, "correct": 4}
    for row in sorted(report.rows, key=lambda r: (order.get(r.status, 9), r.checked_file, r.fact_type)):
        cls = html.escape(row.status)
        checked_ev = row.checked_evidence
        ref_ev = row.reference_evidence
        rows_html.append(f"""
        <tr class="{cls}">
          <td><b>{html.escape(STATUS_LABELS.get(row.status, row.status))}</b></td>
          <td>{html.escape(FACT_LABELS.get(row.fact_type, row.fact_type))}</td>
          <td>{html.escape(row.checked_file)}</td>
          <td>{html.escape(one_line(row.checked_value, 260))}</td>
          <td>{html.escape(one_line(row.reference_value, 260))}</td>
          <td>{row.score}%</td>
          <td>
            <div><b>Проверяемый:</b> {html.escape(_evidence_text(checked_ev))}</div>
            <div><b>Эталон:</b> {html.escape(_evidence_text(ref_ev))}</div>
          </td>
          <td>{html.escape(row.comment)}</td>
        </tr>
        """)


    art_html = []
    art_labels = {"found": "Найдено", "missing": "Не найдено", "needs_review": "Нужна проверка"}
    for req in report.artifact_requirements:
        art_html.append(f"""<tr class="{html.escape(req.status)}"><td><b>{html.escape(art_labels.get(req.status, req.status))}</b></td><td>{html.escape(req.title)}</td><td>{html.escape('; '.join(req.found_files) or '—')}</td><td>{html.escape(req.expected)}</td><td>{html.escape(req.comment)}</td></tr>""")

    addr_html = []
    for row in report.address_consistency:
        addr_html.append(f"""<tr class="{html.escape(row.status)}"><td><b>{html.escape(STATUS_LABELS.get(row.status, row.status))}</b></td><td>{html.escape(row.left_file)}</td><td>{html.escape(one_line(row.left_address, 260))}</td><td>{html.escape(one_line(row.right_address, 260))}</td><td>{row.score}%</td><td>{html.escape(row.comment)}</td></tr>""")

    docs_html = []
    for doc in report.documents:
        facts = ", ".join(f"{html.escape(FACT_LABELS.get(k, k))}: {len(v)}" for k, v in doc.facts.items() if v) or "не найдено"
        docs_html.append(f"<tr><td>{html.escape(doc.filename)}</td><td>{html.escape(doc.role)}</td><td>{doc.stats.get('total_text_chars', 0)}</td><td>{html.escape(facts)}</td></tr>")

    return f"""<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>Отчет проверки ВНИИМТ</title>
<style>
:root {{ --ink:#10232b; --line:#dfe7ea; --brand:#0f4a55; --ok:#d9ead3; --warn:#fff2cc; --bad:#f4cccc; --miss:#fce5cd; }}
body {{ margin:0; font-family:Inter, Arial, sans-serif; color:var(--ink); background:#f6f8f8; }}
.page {{ max-width:1480px; margin:0 auto; padding:28px; }}
.card {{ background:#fff; border:1px solid var(--line); border-radius:22px; padding:22px; margin:18px 0; box-shadow:0 12px 32px rgba(16,35,43,.08); }}
h1 {{ margin:0 0 12px; }}
.actions a {{ display:inline-block; padding:12px 16px; border-radius:999px; background:var(--brand); color:white; text-decoration:none; margin-right:10px; font-weight:800; }}
table {{ border-collapse:collapse; width:100%; font-size:13px; }}
th,td {{ border:1px solid var(--line); padding:8px; vertical-align:top; }}
th {{ background:#e9eef3; text-align:left; }}
.correct td:first-child {{ background:var(--ok); }}
.possible_match td:first-child {{ background:var(--warn); }}
.mismatch td:first-child {{ background:var(--bad); }}
.missing_in_reference td:first-child,.missing_in_document td:first-child {{ background:var(--miss); }}
.found td:first-child {{ background:var(--ok); }}
.missing td:first-child {{ background:var(--bad); }}
.needs_review td:first-child {{ background:var(--warn); }}
.badge {{ display:inline-block; background:#eef4f5; border-radius:999px; padding:7px 11px; margin:4px; }}
</style>
</head>
<body>
<div class="page">
  <div class="card">
    <h1>Отчет первичной проверки документации для ВНИИМТ</h1>
    <p><b>Проект:</b> {html.escape(report.project_name)}</p>
    <p><b>Эталонный документ:</b> {html.escape(report.reference_file)}</p>
    <p><b>Итог:</b> {html.escape(str(report.summary.get('human_status', '')))}</p>
    <div>
      {''.join(f'<span class="badge">{html.escape(STATUS_LABELS.get(k,k))}: {v}</span>' for k,v in status_counts.items())}
    </div>
    <p class="actions">
      <a href="/api/projects/{report.project_id}/download/docx">Скачать DOCX</a>
      <a href="/api/projects/{report.project_id}/download/json">Скачать JSON</a>
    </p>
  </div>
  <div class="card">
    <h2>Что извлек скрипт</h2>
    <table><thead><tr><th>Документ</th><th>Роль</th><th>Символов текста</th><th>Извлечено</th></tr></thead><tbody>{''.join(docs_html)}</tbody></table>
  </div>

  <div class="card">
    <h2>Проверочные артефакты по типу пакета</h2>
    <table><thead><tr><th>Статус</th><th>Правило</th><th>Найденные файлы</th><th>Что требуется</th><th>Комментарий</th></tr></thead><tbody>{''.join(art_html)}</tbody></table>
  </div>
  <div class="card">
    <h2>Отдельная сверка адресов</h2>
    <table><thead><tr><th>Статус</th><th>Документ</th><th>Адрес в документе</th><th>Адрес в эталоне</th><th>Сходство</th><th>Комментарий</th></tr></thead><tbody>{''.join(addr_html) or '<tr><td colspan="6">Адресные пары не найдены или в эталоне нет адреса.</td></tr>'}</tbody></table>
  </div>
  <div class="card">
    <h2>Проверка соответствия</h2>
    <table><thead><tr><th>Статус</th><th>Поле</th><th>Документ</th><th>Значение в документе</th><th>Значение в эталоне</th><th>Сходство</th><th>Откуда взято</th><th>Комментарий</th></tr></thead><tbody>{''.join(rows_html)}</tbody></table>
  </div>
</div>
</body>
</html>"""
