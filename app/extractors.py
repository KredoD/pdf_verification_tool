from __future__ import annotations

import io
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from docx import Document
from PIL import Image

from app.config import settings
from app.models import DocumentReadResult, EvidenceItem
from app.utils import clean_text, find_context_line, normalize_digits, normalize_for_compare, one_line, unique_preserve

try:
    import pytesseract
except Exception:  # OCR is optional at import time
    pytesseract = None

try:
    import openpyxl
except Exception:
    openpyxl = None


# -----------------------------
# Stage 1. Reading documents
# -----------------------------

def _has_ocr() -> bool:
    return bool(settings.ocr_enabled and pytesseract is not None and shutil.which("tesseract"))


def _read_docx_text(path: Path) -> tuple[str, dict[str, Any]]:
    """Extract normal DOCX paragraphs and tables with trace markers."""
    parts: list[str] = []
    stats: dict[str, Any] = {"docx_paragraphs": 0, "docx_tables": 0, "docx_rows": 0}
    doc = Document(str(path))

    for p_index, paragraph in enumerate(doc.paragraphs, start=1):
        text = clean_text(paragraph.text)
        if text:
            stats["docx_paragraphs"] += 1
            parts.append(f"[DOCX_PARAGRAPH_{p_index}]\n{text}")

    for t_index, table in enumerate(doc.tables, start=1):
        stats["docx_tables"] += 1
        rows: list[str] = []
        for r_index, row in enumerate(table.rows, start=1):
            cells = [one_line(cell.text, 1200) for cell in row.cells]
            if any(cells):
                stats["docx_rows"] += 1
                rows.append(f"[ROW_{r_index}] " + " | ".join(cells))
        if rows:
            parts.append(f"[DOCX_TABLE_{t_index}]\n" + "\n".join(rows))
    return clean_text("\n\n".join(parts)), stats


def _read_docx_ocr(path: Path) -> tuple[str, dict[str, Any]]:
    """OCR embedded DOCX images, preserving image names as locations."""
    stats: dict[str, Any] = {"ocr_available": _has_ocr(), "docx_images_total": 0, "docx_images_ocr_used": 0}
    if not _has_ocr():
        return "", stats

    parts: list[str] = []
    with zipfile.ZipFile(path) as zf:
        names = [
            n for n in zf.namelist()
            if n.startswith("word/media/") and n.lower().endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff"))
        ]
        stats["docx_images_total"] = len(names)
        for index, name in enumerate(names[: settings.ocr_max_docx_images], start=1):
            try:
                image = Image.open(io.BytesIO(zf.read(name)))
                text = pytesseract.image_to_string(image, lang=settings.ocr_lang, config="--psm 6")
                text = clean_text(text)
                if text:
                    stats["docx_images_ocr_used"] += 1
                    parts.append(f"[OCR_IMAGE_{index}: {Path(name).name}]\n{text}")
            except Exception as exc:
                parts.append(f"[OCR_IMAGE_{index}: {Path(name).name} ERROR] {exc}")
    return clean_text("\n\n".join(parts)), stats


def _pdf_page_needs_ocr(text: str) -> bool:
    compact = re.sub(r"\s+", "", text or "")
    if len(compact) < 80:
        return True
    alpha = sum(ch.isalpha() for ch in text)
    return alpha / max(1, len(text)) < 0.12


def _read_pdf(path: Path) -> tuple[str, dict[str, Any]]:
    """Read PDF text. OCR only pages that look like scans."""
    stats: dict[str, Any] = {"pdf_pages": 0, "pdf_pages_ocr_used": 0, "ocr_available": _has_ocr()}
    parts: list[str] = []
    with fitz.open(path) as doc:
        stats["pdf_pages"] = doc.page_count
        for page_index, page in enumerate(doc, start=1):
            text = clean_text(page.get_text("text", sort=True) or "")
            page_parts = [f"[PDF_PAGE_{page_index}]"]
            if text:
                page_parts.append(text)
            if _pdf_page_needs_ocr(text) and _has_ocr() and page_index <= settings.ocr_max_pages_per_file:
                try:
                    zoom = settings.ocr_dpi / 72
                    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
                    image = Image.open(io.BytesIO(pix.tobytes("png")))
                    ocr_text = clean_text(pytesseract.image_to_string(image, lang=settings.ocr_lang, config="--psm 6"))
                    if ocr_text and ocr_text not in text:
                        stats["pdf_pages_ocr_used"] += 1
                        page_parts.append("[OCR]\n" + ocr_text)
                except Exception as exc:
                    page_parts.append(f"[OCR_ERROR] {exc}")
            parts.append("\n".join(page_parts))
    return clean_text("\n\n".join(parts)), stats


def _read_xlsx(path: Path) -> tuple[str, dict[str, Any]]:
    stats: dict[str, Any] = {"xlsx_sheets": 0, "xlsx_rows": 0}
    if openpyxl is None:
        return "", {**stats, "error": "openpyxl is not installed"}
    parts: list[str] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    stats["xlsx_sheets"] = len(wb.sheetnames)
    for sheet in wb.worksheets:
        rows = []
        for r_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [one_line(str(c), 600) if c is not None else "" for c in row]
            if any(cells):
                stats["xlsx_rows"] += 1
                rows.append(f"[ROW_{r_index}] " + " | ".join(cells))
        if rows:
            parts.append(f"[XLSX_SHEET: {sheet.title}]\n" + "\n".join(rows))
    return clean_text("\n\n".join(parts)), stats


def read_any_document(path: Path, *, role_hint: str = "") -> DocumentReadResult:
    """Stage A1-A3: read a file and classify its document role."""
    suffix = path.suffix.lower()
    stats: dict[str, Any] = {"filename": path.name, "size_bytes": path.stat().st_size, "extension": suffix}
    text = ""

    try:
        if suffix == ".docx":
            text, s1 = _read_docx_text(path)
            stats.update(s1)
            ocr_text, s2 = _read_docx_ocr(path)
            stats.update(s2)
            if ocr_text:
                text = clean_text("\n\n".join([text, ocr_text]))
        elif suffix == ".pdf":
            text, s = _read_pdf(path)
            stats.update(s)
        elif suffix in {".xlsx", ".xlsm"}:
            text, s = _read_xlsx(path)
            stats.update(s)
        elif suffix in {".txt", ".csv", ".md"}:
            text = path.read_text(encoding="utf-8", errors="ignore")
        else:
            text = path.read_text(encoding="utf-8", errors="ignore")
    except Exception as exc:
        stats["read_error"] = str(exc)
        text = ""

    text = clean_text(text)
    stats["total_text_chars"] = len(text)
    role = role_hint or detect_document_role(path.name, text)
    result = DocumentReadResult(filename=path.name, role=role, text=text, stats=stats)
    result.facts = extract_all_facts(text, path.name)
    return result


# -----------------------------
# Stage 4. Role detection
# -----------------------------

ROLE_KEYWORDS = {
    # Проверочный/master-документ: обычно ЕГРЮЛ, заявление или явно выбранный пользователем эталон.
    "REFERENCE_EGRUL_OR_MASTER": ["егрюл", "выписка", "единый государственный реестр", "master", "эталон"],

    # Заявление важно как отдельный проверочный документ: там часто правильные НКМИ, класс риска, производитель и адреса.
    "REGISTRATION_APPLICATION": ["регистрацион", "заявление", "класс потенциального риска", "код вида", "нкми"],

    # Если есть УПП, то рядом должен быть проверочный артефакт: бизнес-лицензия/разрешение.
    "UPP": ["уполномоченный представитель", "упп", "представител"],
    "BUSINESS_LICENSE": ["бизнес лиценз", "бизнес-лиценз", "лиценз", "license", "business license", "разрешение"],

    # OEM-сценарий: площадка + договор между отечественным производителем и иностранной площадкой.
    "OEM_AGREEMENT": ["oem contract", "oem agreement", "соглашение_oem", "договор oem", "договор между", "контракт"],
    "OEM_SITE": ["oem", "контрактн", "площадка", "foreign site", "иностранная площадка", "производственная площадка"],

    # РФ-производитель: нужен артефакт по праву пользования производством.
    "RF_MANUFACTURER": ["рф производитель", "российский производитель", "отечественный производитель", "производитель рф"],
    "LEASE_AGREEMENT": ["договор аренды", "аренда", "арендодатель", "арендатор"],
    "EGRN_EXTRACT": ["егрн", "единый государственный реестр недвижимости"],
    "OWNERSHIP_DOC": ["собственност", "право собственности", "право пользования"],
    "PRODUCTION_SITE_DOC": ["производственное помещение", "производственная площадка", "адрес производства"],

    "TECHNICAL_DOCUMENT": ["техническ", "втд", "эксплуатац", "инструкция"],
}


def detect_document_role(filename: str, text: str) -> str:
    hay = f"{filename}\n{text[:25000]}".lower()
    best_role = "OTHER_DOCUMENT"
    best_score = 0
    for role, words in ROLE_KEYWORDS.items():
        score = 0
        for w in words:
            if w in filename.lower():
                score += 5
            if w in hay:
                score += 1
        if score > best_score:
            best_role, best_score = role, score
    return best_role


# -----------------------------
# Stage 5. Entity/fact extraction with evidence
# -----------------------------

FACT_PATTERNS: dict[str, list[tuple[str, str]]] = {
    "inn": [(r"\bИНН\s*[:№\-]?\s*(\d{10}|\d{12})\b", "regex:ИНН"), (r"(?<!\d)(\d{10}|\d{12})(?!\d)", "regex:10/12 digits")],
    "ogrn": [(r"\bОГРН(?:ИП)?\s*[:№\-]?\s*(\d{13}|\d{15})\b", "regex:ОГРН"), (r"(?<!\d)(\d{13}|\d{15})(?!\d)", "regex:13/15 digits")],
    "email": [(r"([A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,})", "regex:email")],
    "phone": [(r"((?:\+7|8)[\s(\-]*\d{3}[\s)\-]*\d{3}[\s\-]*\d{2}[\s\-]*\d{2})", "regex:phone")],
    "nkmi_code": [(r"(?:НКМИ|код\s+вида|номенклатурн)[^0-9]{0,60}(\d{5,8})", "regex:НКМИ")],
    "risk_class": [(r"(?:класс(?:а)?\s+(?:потенциального\s+)?риска[^0-9а-я]{0,20})(1|2а|2б|3)\b", "regex:risk_class"), (r"\b(1|2а|2б|3)\s*(?:класс(?:а)?\s+риска)\b", "regex:risk_class")],
    "license": [(r"((?:лицензи[яи]|license|разрешени[ея])[^\n]{0,180})", "regex:license line")],
    "organization": [(r"((?:ООО|АО|ПАО|ЗАО|ОАО|ИП|ФГБУ|ФБУ|ФАУ|НКО|LLC|LTD|CO\.?\s*LTD|INC)[\s\"«„A-Za-zА-Яа-яЕе0-9.,()\-]{2,170})", "regex:organization")],
    "address": [(r"((?:адрес|место нахождения|location|address)[^\n]{0,240})", "regex:address line"), (r"([^\n]{0,180}(?:ул\.|улица|проспект|шоссе|город|г\.|область|край|республика|street|road|avenue)[^\n]{0,180})", "regex:address words")],
}

DIGIT_FACTS = {"inn", "ogrn", "phone", "nkmi_code"}


def _location_from_text_before(text: str, pos: int) -> str:
    """Find nearest marker inserted by the reader before a match."""
    prefix = text[:pos]
    markers = re.findall(r"\[(PDF_PAGE_\d+|DOCX_TABLE_\d+|DOCX_PARAGRAPH_\d+|OCR_IMAGE_\d+:[^\]]+|XLSX_SHEET:[^\]]+)\]", prefix)
    return markers[-1] if markers else "document"


def _normalize_fact(fact_type: str, value: str) -> str:
    if fact_type in DIGIT_FACTS:
        return normalize_digits(value)
    if fact_type == "risk_class":
        return value.lower().replace(" ", "")
    return normalize_for_compare(value)


def extract_all_facts(text: str, source_file: str) -> dict[str, list[EvidenceItem]]:
    """Extract all required fact types and keep evidence for each value."""
    facts: dict[str, list[EvidenceItem]] = {k: [] for k in FACT_PATTERNS}
    seen: set[tuple[str, str]] = set()

    for fact_type, patterns in FACT_PATTERNS.items():
        for pattern, method in patterns:
            for m in re.finditer(pattern, text, flags=re.IGNORECASE | re.MULTILINE):
                value = m.group(1) if m.groups() else m.group(0)
                value = one_line(value, 500)
                normalized = _normalize_fact(fact_type, value)
                if not normalized:
                    continue
                key = (fact_type, normalized)
                if key in seen:
                    continue
                seen.add(key)
                facts[fact_type].append(
                    EvidenceItem(
                        fact_type=fact_type,
                        value=value,
                        normalized_value=normalized,
                        source_file=source_file,
                        location=_location_from_text_before(text, m.start()),
                        quote=find_context_line(text, m.start(), m.end()),
                        extraction_method=method,
                        confidence=95 if method.startswith("regex") else 80,
                    )
                )

    return facts
